from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\Politeknik Statistika STIS\Pembelajaran\Skripsi\SourceCode\SkripsiNabil")
REFERENCE = Path(r"D:\Politeknik Statistika STIS\Pembelajaran\Skripsi\[1] [On going] SLM - RAG\[1] Progres Skripsi\[4] Bab 4\tujuan 1.docx")
SOURCE = ROOT / "Hasil" / "tujuan2" / "Draf Narasi Tujuan 2 - Hasil Fine-Tuning SLM.md"
ASSETS = ROOT / "Hasil" / "tujuan2" / "word_build" / "assets"
OUTPUT = Path(r"D:\Politeknik Statistika STIS\Pembelajaran\Skripsi\[1] [On going] SLM - RAG\[1] Progres Skripsi\[4] Bab 4\Tujuan 2 - Hasil Fine-Tuning SLM.docx")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="B8BEC5", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_repeat_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def strip_markdown(text):
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = text.replace("&lt;", "<").replace("&gt;", ">")
    return text.strip()


def extract_status(text):
    m = re.search(r"\s*\*\*\[([^\]]+)\]\*\*\s*$", text)
    if not m:
        return text, None
    return text[:m.start()].rstrip(), m.group(1)


def add_text_with_emphasis(paragraph, text):
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


doc = Document(str(REFERENCE))
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
available_style_names = [s.name for s in styles]
def style_named(name):
    for candidate in styles:
        if candidate.name == name:
            return candidate
    raise KeyError(name)

if "Normal" not in available_style_names:
    normal = styles.add_style("Normal", WD_STYLE_TYPE.PARAGRAPH)
else:
    normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

for style_name, size in (("Heading 2", 13), ("Heading 3", 12)):
    st = style_named(style_name)
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(8)
    st.paragraph_format.space_after = Pt(4)

if "List Number" not in [s.name for s in styles]:
    list_number = styles.add_style("List Number", WD_STYLE_TYPE.PARAGRAPH)
    list_number.base_style = normal
if "List Bullet" not in [s.name for s in styles]:
    list_bullet = styles.add_style("List Bullet", WD_STYLE_TYPE.PARAGRAPH)
    list_bullet.base_style = normal

if "Subjudul Bab" not in [s.name for s in styles]:
    sub_style = styles.add_style("Subjudul Bab", WD_STYLE_TYPE.PARAGRAPH)
else:
    sub_style = styles["Subjudul Bab"]
sub_style.font.name = "Times New Roman"
sub_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
sub_style.font.size = Pt(12)
sub_style.font.bold = True
sub_style.font.color.rgb = RGBColor(45, 45, 45)
sub_style.paragraph_format.space_before = Pt(6)
sub_style.paragraph_format.space_after = Pt(3)
sub_style.paragraph_format.keep_with_next = True


def add_body(text, status=None, first_indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.4)
    add_text_with_emphasis(p, text)
    if status:
        r = p.add_run(f"  [{status}]")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(83, 103, 120)
    return p


def add_caption(text, above=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3 if above else 6)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.bold = False
    keep_with_next(p, above)
    return p


def add_source_note(text="Sumber: Hasil pengolahan peneliti"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = RGBColor(80, 80, 80)


def add_figure(filename, caption, width=5.9):
    path = ASSETS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    p.paragraph_format.keep_with_next = True
    add_caption(caption, above=False)
    add_source_note()


def add_table(rows, caption, col_widths=None, font_size=9):
    cap = add_caption(caption, above=True)
    if any(caption.startswith(f"Tabel {number} ") for number in ("4.3", "4.4", "4.5", "4.7")):
        cap.paragraph_format.page_break_before = True
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, row in enumerate(rows):
        set_repeat_no_split(table.rows[i])
        if i == 0:
            set_repeat_table_header(table.rows[i])
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if col_widths:
                cell.width = Inches(col_widths[j])
            if i == 0:
                shade_cell(cell, "3F4A56")
            elif i % 2 == 0:
                shade_cell(cell, "F3F5F7")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(strip_markdown(value))
            r.font.name = "Times New Roman"
            r.font.size = Pt(font_size)
            r.bold = i == 0
            if i == 0:
                r.font.color.rgb = RGBColor(255, 255, 255)
    add_source_note()
    return table


def parse_md_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x) for x in cells):
            rows.append(cells)
        i += 1
    return rows, i


def table_spec(header):
    joined = "|".join(header)
    if joined.startswith("Aspek|"):
        return "Tabel 4.3 Ringkasan dataset fine-tuning setelah penyaringan", [3.8, 2.3], 9.5
    if joined.startswith("Parameter|"):
        return "Tabel 4.4 Konfigurasi QLoRA pada kedua model", [2.0, 2.05, 2.05], 8.5
    if joined.startswith("Model|Kondisi|"):
        return "Tabel 4.6 Hasil evaluasi otomatis model dasar dan fine-tuned", [1.25, 1.0, 1.0, 0.85, 0.85, 1.15], 8.5
    if joined.startswith("Sub-tugas|"):
        return "Tabel 4.7 Perubahan Token F1 menurut jenis pertanyaan", [1.55, 1.35, 0.75, 1.35, 0.75], 8.0
    return "Tabel ringkasan hasil", None, 9


doc.add_paragraph("4.2  Hasil Fine-Tuning Small Language Model", style=style_named("Heading 2"))

note = doc.add_table(rows=1, cols=1)
note.alignment = WD_TABLE_ALIGNMENT.CENTER
note.autofit = False
note.cell(0, 0).width = Inches(6.1)
shade_cell(note.cell(0, 0), "EEF3F7")
set_cell_margins(note.cell(0, 0), top=100, start=120, bottom=100, end=120)
np = note.cell(0, 0).paragraphs[0]
np.paragraph_format.space_after = Pt(0)
np.paragraph_format.line_spacing = 1.0
nr = np.add_run("Keterangan status evidensi: ")
nr.bold = True
nr.font.name = "Times New Roman"
nr.font.size = Pt(9)
nr = np.add_run("[HASIL] berasal langsung dari kode atau keluaran eksperimen; [INTERPRETASI] merupakan penafsiran atas hasil; [ASUMSI] belum diuji langsung; [PERLU DILENGKAPI] menandai artefak atau analisis yang belum tersedia.")
nr.font.name = "Times New Roman"
nr.font.size = Pt(9)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
start = next(i for i, line in enumerate(lines) if line.strip() == "## 4.2 Hasil Fine-Tuning Small Language Model") + 1
i = start
skip_recommendations = False
inserted_training_table = False
inserted_winloss = False

while i < len(lines):
    raw = lines[i].rstrip()
    line = raw.strip()
    if line == "## Rekomendasi urutan tabel dan gambar di Bab IV":
        break
    if not line:
        i += 1
        continue
    if line.startswith("Visual yang paling relevan") or line.startswith("Visual utama untuk"):
        i += 1
        continue
    if strip_markdown(line) == "Tabel ringkasan dataset akhir":
        i += 1
        continue
    if line.startswith("|"):
        rows, next_i = parse_md_table(lines, i)
        caption, widths, fs = table_spec(rows[0])
        add_table(rows, caption, widths, fs)
        header = "|".join(rows[0])
        if header.startswith("Aspek|"):
            add_figure("gambar_dataset_funnel.png", "Gambar 4.9 Alur pembentukan dataset fine-tuning", 5.8)
            add_figure("gambar_cakupan_korpus.png", "Gambar 4.10 Cakupan dokumen korpus dalam dataset sintetis", 5.8)
        elif header.startswith("Model|Kondisi|"):
            add_figure("gambar_metrik_otomatis.png", "Gambar 4.12 Perbandingan metrik otomatis sebelum dan sesudah fine-tuning", 5.9)
        elif header.startswith("Sub-tugas|"):
            add_figure("gambar_subtask.png", "Gambar 4.13 Peningkatan Token F1 menurut jenis pertanyaan", 5.9)
        i = next_i
        continue
    if line.startswith("### "):
        title = strip_markdown(line[4:])
        if title == "Sintesis Tujuan 2":
            title = "4.2.5 Sintesis Hasil dan Batas Kesimpulan"
        doc.add_paragraph(title, style=style_named("Heading 3"))
        i += 1
        continue
    if line.startswith("#### "):
        doc.add_paragraph(strip_markdown(line[5:]), style="Subjudul Bab")
        i += 1
        continue
    if line.startswith("## "):
        title = strip_markdown(line[3:])
        if title == "Sintesis Tujuan 2":
            doc.add_paragraph("4.2.5 Sintesis Hasil dan Batas Kesimpulan", style=style_named("Heading 3"))
        elif title == "Daftar asumsi dan interpretasi yang belum diuji langsung":
            doc.add_paragraph("4.2.6 Asumsi dan Interpretasi yang Belum Diuji Langsung", style=style_named("Heading 3"))
        elif title == "Daftar temuan yang masih kurang atau perlu diperbaiki":
            doc.add_paragraph("4.2.7 Temuan yang Masih Perlu Dilengkapi", style=style_named("Heading 3"))
        i += 1
        continue
    if line.startswith(">"):
        quote = line.lstrip("> ")
        quote, status = extract_status(quote)
        p = add_body(quote, status)
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.right_indent = Inches(0.25)
        for run in p.runs:
            run.italic = True
        i += 1
        continue
    mlist = re.match(r"^(\d+)\.\s+(.*)$", line)
    if mlist:
        txt, status = extract_status(mlist.group(2))
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        add_text_with_emphasis(p, txt)
        for rr in p.runs:
            rr.font.size = Pt(10)
        if status:
            rr = p.add_run(f" [{status}]")
            rr.italic = True
            rr.font.size = Pt(9)
            rr.font.color.rgb = RGBColor(83, 103, 120)
        i += 1
        continue
    if line.startswith("- "):
        txt, status = extract_status(line[2:])
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        add_text_with_emphasis(p, txt)
        if status:
            rr = p.add_run(f" [{status}]")
            rr.italic = True
            rr.font.size = Pt(9)
        i += 1
        continue

    text, status = extract_status(line)
    add_body(text, status)

    if text.startswith("Waktu Gemma-2 dan Llama-3.2 tidak boleh dibandingkan") and not inserted_training_table:
        rows = [
            ["Model", "Parameter dilatih", "Waktu latih", "Training loss", "Validation loss terakhir"],
            ["Gemma-2-2B", "20.766.720 (0,79%)", "4,00 jam", "0,4376", "0,4405"],
            ["Llama-3.2-3B", "24.313.856 (0,75%)", "5,75 jam", "0,3609", "0,3404"],
        ]
        add_table(rows, "Tabel 4.5 Ringkasan hasil pelatihan kedua model", [1.25, 1.4, 1.0, 1.15, 1.35], 8.5)
        add_figure("gambar_kurva_loss.png", "Gambar 4.11 Kurva training loss dan validation loss", 5.9)
        inserted_training_table = True
    if text.startswith("Peningkatan tidak hanya terlihat pada rerata") and not inserted_winloss:
        add_figure("gambar_menang_kalah.png", "Gambar 4.14 Proporsi butir yang membaik, seri, dan memburuk setelah fine-tuning", 5.9)
        inserted_winloss = True
    if text.startswith("Sampai pemeriksaan ini dilakukan, berkas"):
        rows = [
            ["Komponen evaluasi manusia", "Status saat penyusunan", "Implikasi"],
            ["Skor lima evaluator", "Belum tersedia", "Rerata fluency, factual correctness, dan completeness belum dapat dihitung"],
            ["Uji beda base–fine-tuned", "Belum tersedia", "Tidak ada dasar untuk menyatakan peningkatan menurut evaluator"],
            ["Krippendorff's alpha ordinal", "Belum dihitung", "Reliabilitas antarevaluator belum diketahui"],
            ["Penandaan acuan meragukan", "Parsing perlu diperbaiki", "Nilai 'ya' saat ini berubah menjadi NaN"],
        ]
        add_table(rows, "Tabel 4.8 Status ketersediaan hasil evaluasi manusia", [1.7, 1.35, 3.05], 8.5)
    i += 1


for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.font.name is None:
            run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name or "Times New Roman")

doc.core_properties.title = "Tujuan 2 - Hasil Fine-Tuning Small Language Model"
doc.core_properties.subject = "Bab IV hasil dan pembahasan tujuan penelitian kedua"
doc.core_properties.author = "Nabil"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT)
